import tkinter as tk
from tkinter import ttk, messagebox
from datetime import date, datetime
import os, re, openpyxl
from openpyxl.styles import Font, PatternFill

# SUNMIT Travel - Made Simple | Python + Tkinter + Excel
APP_TITLE = "SUNMIT Travel | Made Simple"
FILE_NAME = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sunmit_travel_bookings.xlsx")
BG, NAVY, BLUE, CYAN, GREEN, RED, WHITE = "#F3F7FC", "#10213D", "#159ED9", "#39C7EF", "#16A765", "#E6535D", "#FFFFFF"
TEXT, MUTED, LINE = "#172B4D", "#7183A1", "#DCE6F2"
CITIES = ["Pune", "Mumbai", "Nashik", "Shirpur", "Dhule", "Indore", "Surat", "Ahmedabad", "Aurangabad", "Nagpur"]
# Sample operator / fare data. Change these values to suit your project.
BUSES = {
    "SUNMIT Express - AC Seater": {"type": "AC Seater", "fare": 650},
    "Royal Roadlines - Non-AC Seater": {"type": "Non-AC Seater", "fare": 350},
    "BlueLine Travels - AC Sleeper": {"type": "AC Sleeper", "fare": 950},
    "Maharaja Coaches - Volvo Sleeper": {"type": "Volvo AC Sleeper", "fare": 1350},
    "Night Rider - Non-AC Sleeper": {"type": "Non-AC Sleeper", "fare": 600},
}
HEADERS = ["Booking ID", "Booked At", "Journey Date", "From City", "To City", "Bus Name", "Bus Type", "Passenger Name", "Contact Number", "Address", "Seat Numbers", "Passenger Genders", "Couple Booking", "Passengers", "Fare Per Seat", "Total Amount", "Status"]


def setup_excel():
    if os.path.exists(FILE_NAME):
        wb = openpyxl.load_workbook(FILE_NAME)
        if "Bookings" in wb.sheetnames:
            old = wb["Bookings"]
            # Preserve an older app workbook instead of silently mixing different column formats.
            if old.max_row > 1 and old.cell(1, 1).value != "Booking ID":
                if "PreviousBookings" not in wb.sheetnames:
                    backup = wb.create_sheet("PreviousBookings")
                    for row in old.iter_rows(values_only=True): backup.append(list(row))
                wb.remove(old)
                ws = wb.create_sheet("Bookings", 0)
                ws.append(HEADERS)
            else:
                ws = old
                if ws.max_row == 1 and ws.cell(1, 1).value != "Booking ID":
                    ws.delete_rows(1, ws.max_row)
                    ws.append(HEADERS)
                elif [ws.cell(1, c).value for c in range(1, len(HEADERS)+1)] != HEADERS:
                    if "PreviousBookings" not in wb.sheetnames:
                        backup = wb.create_sheet("PreviousBookings")
                        for row in ws.iter_rows(values_only=True): backup.append(list(row))
                    wb.remove(ws)
                    ws = wb.create_sheet("Bookings", 0)
                    ws.append(HEADERS)
        else:
            ws = wb.create_sheet("Bookings", 0); ws.append(HEADERS)
    else:
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Bookings"; ws.append(HEADERS)
    ws = wb["Bookings"]
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="16365C")
    wb.save(FILE_NAME)
    return wb

wb = setup_excel()

def rows():
    return [list(r) for r in wb["Bookings"].iter_rows(min_row=2, values_only=True) if r and r[0]]
def active(r): return str(r[16] or "Confirmed") == "Confirmed"
def money(n): return f"₹{float(n or 0):,.0f}"
def is_name_ok(s): return bool(s.strip()) and all(ch.isalpha() or ch.isspace() or ch in ".-'" for ch in s)
def seat_list(bus_name):
    # Realistic simple layouts: seater is 2+2; sleeper is 2 lower/upper berths + 1 single berth.
    if "Sleeper" in BUSES[bus_name]["type"]:
        return [(f"L{r}", f"U{r}", f"S{r}") for r in range(1, 9)]
    return [(f"A{r}", f"B{r}", f"C{r}", f"D{r}") for r in range(1, 11)]
def neighbours(seat, bus_type):
    m = re.search(r"(\d+)$", seat)
    if not m or "Sleeper" in bus_type: return set()
    r = m.group(1)
    groups = [[f"A{r}", f"B{r}"], [f"C{r}", f"D{r}"]]
    for g in groups:
        if seat in g: return set(g) - {seat}
    return set()

root = tk.Tk(); root.title(APP_TITLE); root.geometry("1390x860"); root.minsize(1120, 720); root.configure(bg=BG)
style = ttk.Style(); style.theme_use("clam")
style.configure("Treeview", background=WHITE, foreground=TEXT, fieldbackground=WHITE, rowheight=29, font=("Segoe UI", 9))
style.configure("Treeview.Heading", background="#E7EFF8", foreground=NAVY, font=("Segoe UI", 9, "bold"))
style.configure("TCombobox", padding=5)
sidebar = tk.Frame(root, bg=NAVY, width=225); sidebar.pack(side="left", fill="y"); sidebar.pack_propagate(False)
tk.Label(sidebar, text="🚌 SUNMIT Travel", bg=NAVY, fg=CYAN, font=("Segoe UI", 17, "bold")).pack(pady=(25, 4), padx=8)
tk.Label(sidebar, text="M A D E   S I M P L E", bg=NAVY, fg=WHITE, font=("Segoe UI", 9, "bold")).pack(pady=(0, 32))
main = tk.Frame(root, bg=BG); main.pack(side="left", fill="both", expand=True)
header = tk.Frame(main, bg=WHITE, height=72, highlightbackground=LINE, highlightthickness=1); header.pack(fill="x"); header.pack_propagate(False)
title = tk.Label(header, text="Book a Ticket", bg=WHITE, fg=NAVY, font=("Segoe UI", 20, "bold")); title.pack(side="left", padx=24)
tk.Label(header, text="● SYSTEM READY  |  ADMIN", bg=WHITE, fg=MUTED, font=("Segoe UI", 9, "bold")).pack(side="right", padx=22)
content = tk.Frame(main, bg=BG); content.pack(fill="both", expand=True, padx=16, pady=14)
nav = {}
def clear_page(t):
    for w in content.winfo_children(): w.destroy()
    title.config(text=t)
def card(parent, heading):
    f=tk.Frame(parent,bg=WHITE,highlightbackground=LINE,highlightthickness=1)
    tk.Label(f,text=heading,bg=WHITE,fg=NAVY,font=("Segoe UI",11,"bold")).pack(anchor="w",padx=12,pady=11)
    tk.Frame(f,bg=LINE,height=1).pack(fill="x",padx=12)
    return f
def button(parent, text, cmd, color=BLUE, fg=WHITE):
    return tk.Button(parent,text=text,command=cmd,bg=color,fg=fg,activebackground=color,activeforeground=fg,relief="flat",bd=0,cursor="hand2",font=("Segoe UI",9,"bold"),padx=12,pady=9)
def stats():
    active_rows=[r for r in rows() if active(r)]; today=date.today().isoformat()
    return len(active_rows), sum(1 for r in active_rows if str(r[1]).startswith(today)), sum(float(r[15] or 0) for r in active_rows)
def add_stats():
    box=tk.Frame(content,bg=BG); box.pack(fill="x",pady=(0,10))
    a,b,c=stats()
    for label,val,col in [("CONFIRMED BOOKINGS",a,BLUE),("BOOKINGS TODAY",b,GREEN),("ACTIVE TICKET REVENUE",money(c),"#7357D9")]:
        f=tk.Frame(box,bg=WHITE,highlightbackground=LINE,highlightthickness=1); f.pack(side="left",fill="x",expand=True,padx=4)
        tk.Label(f,text=label,bg=WHITE,fg=MUTED,font=("Segoe UI",8,"bold")).pack(anchor="w",padx=13,pady=(10,2))
        tk.Label(f,text=str(val),bg=WHITE,fg=col,font=("Segoe UI",20,"bold")).pack(anchor="w",padx=13,pady=(0,10))
def set_nav(key):
    for k,v in nav.items(): v.config(bg="#10486A" if k==key else NAVY,fg=WHITE if k==key else "#B8C8DE")
def add_nav(key,label,cmd):
    b=tk.Button(sidebar,text=label,command=cmd,anchor="w",padx=17,pady=13,bg=NAVY,fg="#B8C8DE",activebackground="#10486A",activeforeground=WHITE,relief="flat",bd=0,cursor="hand2",font=("Segoe UI",10,"bold")); b.pack(fill="x",padx=8,pady=3); nav[key]=b

# Booking page
def show_booking():
    clear_page("Book a Ticket"); set_nav("Book Ticket"); add_stats()
    body=tk.Frame(content,bg=BG); body.pack(fill="both",expand=True)
    for c,w in enumerate((4,5,3)): body.grid_columnconfigure(c,weight=w,uniform="cols")
    body.grid_rowconfigure(0,weight=1)
    left=card(body,"Passenger & Journey Details"); left.grid(row=0,column=0,sticky="nsew",padx=(0,5))
    mid=card(body,"Choose Seats  |  Front of Bus at Top"); mid.grid(row=0,column=1,sticky="nsew",padx=5)
    right=card(body,"Booking Summary"); right.grid(row=0,column=2,sticky="nsew",padx=(5,0))
    form=tk.Frame(left,bg=WHITE); form.pack(fill="both",expand=True,padx=12,pady=8); form.grid_columnconfigure(0,weight=1)
    name=tk.StringVar(); contact=tk.StringVar(); address=tk.StringVar(); source=tk.StringVar(value="Pune"); dest=tk.StringVar(value="Mumbai"); jdate=tk.StringVar(value=date.today().isoformat()); bus=tk.StringVar(value=list(BUSES)[0]); count=tk.StringVar(value="1"); gender=tk.StringVar(value="Female"); couple=tk.BooleanVar(value=False)
    def label(text,row): tk.Label(form,text=text,bg=WHITE,fg=MUTED,font=("Segoe UI",8,"bold")).grid(row=row,column=0,sticky="w",pady=(5,2))
    def entry(text,var,row,validate=None):
        label(text,row); e=tk.Entry(form,textvariable=var,bg="#FBFDFF",fg=TEXT,relief="solid",bd=1,font=("Segoe UI",9),validate="key" if validate else "none",validatecommand=validate if validate else None); e.grid(row=row+1,column=0,sticky="ew",ipady=5); return e
    def combo(text,var,values,row):
        label(text,row); c=ttk.Combobox(form,textvariable=var,values=values,state="readonly",font=("Segoe UI",9)); c.grid(row=row+1,column=0,sticky="ew",ipady=3); return c
    name_ok=(root.register(lambda s: all(ch.isalpha() or ch.isspace() or ch in ".-'" for ch in s)),"%P")
    contact_ok=(root.register(lambda s: s.isdigit() or s==""),"%P")
    entry("PASSENGER NAME * (letters only)",name,0,name_ok); entry("MOBILE NUMBER * (digits only)",contact,2,contact_ok); entry("ADDRESS",address,4)
    combo("FROM CITY *",source,CITIES,6); combo("TO CITY *",dest,[c for c in CITIES if c!="Pune"],8); entry("JOURNEY DATE * (YYYY-MM-DD)",jdate,10)
    combo("BUS NAME / OPERATOR *",bus,list(BUSES),12); combo("PASSENGER COUNT",count,["1","2","3","4","5","6"],14); combo("PASSENGER GENDER",gender,["Female","Male","Prefer not to say"],16)
    tk.Checkbutton(form,text="Couple booking (husband/wife or partners)",variable=couple,bg=WHITE,fg=TEXT,activebackground=WHITE,font=("Segoe UI",8)).grid(row=18,column=0,sticky="w",pady=7)
    # Seat layout area
    seat_area=tk.Frame(mid,bg=WHITE); seat_area.pack(fill="both",expand=True,padx=10,pady=8)
    tk.Label(seat_area,text="FRONT  •  DRIVER",bg="#EDF3FA",fg=MUTED,font=("Segoe UI",8,"bold"),pady=6).pack(fill="x")
    grid=tk.Frame(seat_area,bg=WHITE); grid.pack(pady=8)
    seats={}; selected=[]
    def occupied():
        result={}
        for r in rows():
            if str(r[2])==jdate.get() and str(r[3])==source.get() and str(r[4])==dest.get() and str(r[5])==bus.get() and active(r):
                ss=[s.strip() for s in str(r[10] or "").split(",") if s.strip()]; gg=[g.strip() for g in str(r[11] or "").split(",") if g.strip()]
                for i,s in enumerate(ss): result[s]=gg[i] if i<len(gg) else "Unknown"
        return result
    def seat_neighbours(s):
        m=re.search(r"(\d+)$",s)
        if not m or "Sleeper" in BUSES[bus.get()]["type"]: return set()
        n=m.group(1)
        return ({f"A{n}",f"B{n}"}-{s}) if s[0] in "AB" else ({f"C{n}",f"D{n}"}-{s})
    def refresh_seats():
        taken=occupied()
        for s,w in seats.items():
            if s in taken: w.config(bg=RED,fg=WHITE,state="disabled")
            elif s in selected: w.config(bg=CYAN,fg=NAVY,state="normal")
            else: w.config(bg="#E8EEF5",fg=TEXT,state="normal")
    def select_seat(s):
        if s in selected: selected.remove(s)
        elif len(selected)>=int(count.get()): messagebox.showwarning("Seat selection",f"Select only {count.get()} seat(s)."); return
        else: selected.append(s)
        refresh_seats(); update_summary()
    def rebuild_seats():
        for w in grid.winfo_children(): w.destroy()
        seats.clear()
        if "Sleeper" in BUSES[bus.get()]["type"]:
            tk.Label(grid,text="LOWER / UPPER",bg=WHITE,fg=MUTED,font=("Segoe UI",7,"bold")).grid(row=0,column=0,columnspan=2)
            tk.Label(grid,text="SINGLE BERTH",bg=WHITE,fg=MUTED,font=("Segoe UI",7,"bold")).grid(row=0,column=4)
            for r in range(1,9):
                for s,c in [(f"L{r}",0),(f"U{r}",1),(f"S{r}",4)]:
                    w=tk.Button(grid,text=s,width=5,font=("Segoe UI",8,"bold"),bd=0,command=lambda x=s:select_seat(x)); w.grid(row=r,column=c,padx=4,pady=3); seats[s]=w
        else:
            for r in range(1,11):
                tk.Label(grid,text=str(r),bg=WHITE,fg=MUTED,font=("Segoe UI",8)).grid(row=r,column=0,padx=3)
                for s,c in [(f"A{r}",1),(f"B{r}",2),(f"C{r}",4),(f"D{r}",5)]:
                    w=tk.Button(grid,text=s,width=5,font=("Segoe UI",8,"bold"),bd=0,command=lambda x=s:select_seat(x)); w.grid(row=r,column=c,padx=3,pady=3); seats[s]=w
        refresh_seats()
    tk.Label(seat_area,text="Available  •  Selected  •  Booked",bg=WHITE,fg=MUTED,font=("Segoe UI",8)).pack(pady=4)
    summary=tk.Frame(right,bg=WHITE); summary.pack(fill="both",expand=True,padx=12,pady=8); values={}
    for label_text,key in [("Route","route"),("Journey date","date"),("Bus","bus"),("Bus type","type"),("Seats","seats"),("Passengers","count"),("Fare / seat","fare")]:
        line=tk.Frame(summary,bg=WHITE); line.pack(fill="x",pady=7); tk.Label(line,text=label_text,bg=WHITE,fg=MUTED,font=("Segoe UI",8)).pack(side="left")
        v=tk.Label(line,text="—",bg=WHITE,fg=TEXT,font=("Segoe UI",8,"bold"),wraplength=130,justify="right"); v.pack(side="right"); values[key]=v
    total=tk.Label(summary,text="TOTAL: ₹650",bg="#EDF6FF",fg=NAVY,font=("Segoe UI",14,"bold"),pady=12); total.pack(fill="x",pady=8)
    def update_summary(*args):
        info=BUSES[bus.get()]; fare=info["fare"]
        vals={"route":f"{source.get()} → {dest.get()}","date":jdate.get(),"bus":bus.get(),"type":info["type"],"seats":", ".join(selected) or "Choose seats","count":count.get(),"fare":money(fare)}
        for k,v in vals.items(): values[k].config(text=v)
        total.config(text=f"TOTAL: {money(fare*int(count.get()))}")
    def save_booking():
        nm=name.get().strip(); mob=contact.get().strip()
        if not is_name_ok(nm): messagebox.showerror("Invalid name","Name must contain letters only (spaces, dot, apostrophe and hyphen are allowed)."); return
        if not mob.isdigit() or not 10<=len(mob)<=13: messagebox.showerror("Invalid number","Mobile number must contain 10 to 13 digits only."); return
        if source.get()==dest.get(): messagebox.showerror("Invalid route","From and To cities must be different."); return
        try:
            d=datetime.strptime(jdate.get(),"%Y-%m-%d").date()
            if d<date.today(): raise ValueError
        except ValueError: messagebox.showerror("Invalid date","Enter a valid future date in YYYY-MM-DD format."); return
        n=int(count.get())
        if len(selected)!=n: messagebox.showerror("Select seats",f"Select exactly {n} seat(s)."); return
        taken=occupied()
        if any(s in taken for s in selected): messagebox.showerror("Seat conflict","A selected seat is already booked."); refresh_seats(); return
        # Same-gender neighbour preference; couple booking is the explicit exception.
        my_gender=gender.get()
        if not couple.get() and my_gender in ("Male","Female"):
            for s in selected:
                for other in seat_neighbours(s):
                    if taken.get(other) in ("Male","Female") and taken[other]!=my_gender:
                        messagebox.showerror("Seat gender preference",f"Seat {s} is next to a {taken[other].lower()} passenger. Choose another seat. Couple booking should only be used by people travelling together."); return
        if couple.get() and n<2: messagebox.showerror("Couple booking","Select at least 2 passengers for couple booking."); return
        if couple.get() and n%2: messagebox.showerror("Couple booking","For this demo, couple booking requires an even number of passengers."); return
        if couple.get():
            # Each couple must select neighbouring seats (e.g. A1+B1 or C1+D1).
            remaining = set(selected)
            while remaining:
                first = remaining.pop()
                pair = neighbours(first, BUSES[bus.get()]["type"]) & remaining
                if not pair:
                    messagebox.showerror("Couple seats", "Couple seats must be side-by-side on the same side of the aisle.")
                    return
                remaining.remove(next(iter(pair)))
        genders=[my_gender if not couple.get() or i%2==0 else "Partner" for i in range(n)]
        info=BUSES[bus.get()]; fare=info["fare"]; amount=fare*n; bid="SUN"+datetime.now().strftime("%y%m%d%H%M%S")
        ws=wb["Bookings"]; ws.append([bid,datetime.now().strftime("%Y-%m-%d %H:%M:%S"),jdate.get(),source.get(),dest.get(),bus.get(),info["type"],nm,mob,address.get().strip(),", ".join(selected),", ".join(genders),"Yes" if couple.get() else "No",n,fare,amount,"Confirmed"])
        try: wb.save(FILE_NAME)
        except PermissionError:
            ws.delete_rows(ws.max_row,1); messagebox.showerror("Excel file open","Close sunmit_travel_bookings.xlsx in Excel, then book again."); return
        messagebox.showinfo("Booking confirmed",f"SUNMIT Travel booking successful!\n\nBooking ID: {bid}\nPassenger: {nm}\nRoute: {source.get()} → {dest.get()}\nBus: {bus.get()}\nSeats: {', '.join(selected)}\nTotal: {money(amount)}\n\nRevenue has been updated.")
        show_booking()
    button(summary,"✓  CONFIRM BOOKING",save_booking,GREEN).pack(fill="x",pady=(4,7)); button(summary,"↻  RESET FORM",show_booking,"#E8EEF5",TEXT).pack(fill="x")
    def route_update(*_):
        if source.get()==dest.get(): dest.set(next((c for c in CITIES if c!=source.get()),"Mumbai"))
        refresh_seats(); update_summary()
    source.trace_add("write",route_update); dest.trace_add("write",lambda *_:(refresh_seats(),update_summary())); jdate.trace_add("write",lambda *_:(refresh_seats(),update_summary()))
    bus.trace_add("write",lambda *_:(rebuild_seats(),update_summary())); count.trace_add("write",lambda *_:(update_summary(), trim_selected()))
    def trim_selected():
        while len(selected)>int(count.get()): selected.pop()
        refresh_seats(); update_summary()
    rebuild_seats(); update_summary()

def show_bookings():
    clear_page("Bookings"); set_nav("Bookings")
    top=tk.Frame(content,bg=BG); top.pack(fill="x",pady=(0,10)); query=tk.StringVar()
    tk.Entry(top,textvariable=query,width=32,font=("Segoe UI",10),relief="solid",bd=1).pack(side="left",ipady=6)
    panel=card(content,"Booking Management"); panel.pack(fill="both",expand=True)
    cols=("id","name","contact","route","date","bus","seats","amount","status"); tree=ttk.Treeview(panel,columns=cols,show="headings")
    for c,label,w in [("id","BOOKING ID",160),("name","PASSENGER",130),("contact","CONTACT",110),("route","ROUTE",145),("date","DATE",100),("bus","BUS NAME",170),("seats","SEATS",100),("amount","TOTAL",90),("status","STATUS",90)]: tree.heading(c,text=label); tree.column(c,width=w)
    tree.pack(fill="both",expand=True,padx=12,pady=12)
    def refresh(*_):
        tree.delete(*tree.get_children())
        for r in reversed(rows()):
            if query.get().lower() not in " ".join(str(x or "") for x in r).lower(): continue
            tree.insert("","end",values=(r[0],r[7],r[8],f"{r[3]} → {r[4]}",r[2],r[5],r[10],money(r[15]),r[16]))
    def cancel():
        sel=tree.selection()
        if not sel: messagebox.showwarning("Select booking","Select a booking first."); return
        bid=tree.item(sel[0],"values")[0]
        if not messagebox.askyesno("Cancel booking",f"Cancel {bid}?"): return
        ws=wb["Bookings"]
        for i in range(2,ws.max_row+1):
            if str(ws.cell(i,1).value)==bid and ws.cell(i,17).value!="Cancelled": ws.cell(i,17).value="Cancelled"; break
        wb.save(FILE_NAME); refresh(); messagebox.showinfo("Cancelled","Booking cancelled. Revenue recalculated automatically.")
    button(top,"Search",refresh).pack(side="left",padx=8); button(top,"Cancel selected",cancel,RED).pack(side="right")
    query.trace_add("write",refresh); refresh()
def show_dashboard():
    clear_page("Dashboard"); set_nav("Dashboard"); add_stats(); panel=card(content,"Recent Bookings"); panel.pack(fill="both",expand=True)
    tree=ttk.Treeview(panel,columns=("id","name","route","date","bus","amount"),show="headings")
    for c,t,w in [("id","BOOKING ID",165),("name","PASSENGER",140),("route","ROUTE",150),("date","JOURNEY DATE",110),("bus","BUS OPERATOR",180),("amount","TOTAL",90)]: tree.heading(c,text=t); tree.column(c,width=w)
    for r in reversed(rows()[-15:]): tree.insert("","end",values=(r[0],r[7],f"{r[3]} → {r[4]}",r[2],r[5],money(r[15])))
    tree.pack(fill="both",expand=True,padx=12,pady=12)
def show_reports():
    clear_page("Reports"); set_nav("Reports"); add_stats(); panel=card(content,"Revenue by Bus Operator"); panel.pack(fill="both",expand=True)
    summary={}
    for r in rows():
        if active(r): summary[r[5]]=summary.get(r[5],0)+float(r[15] or 0)
    for name,value in sorted(summary.items()): tk.Label(panel,text=f"{name}     Revenue: {money(value)}",bg=WHITE,fg=TEXT,font=("Segoe UI",10)).pack(anchor="w",padx=16,pady=7)
def show_settings():
    clear_page("Settings"); set_nav("Settings"); panel=card(content,"System Information"); panel.pack(fill="x")
    tk.Label(panel,text=f"SUNMIT Travel — Made Simple\nData file: {FILE_NAME}\n\nSample bus names and fares are configured in the BUSES dictionary near the top of this code.\nThis is a local Excel demo, not connected to live bus operators or a payment gateway.",bg=WHITE,fg=TEXT,justify="left",font=("Segoe UI",10)).pack(anchor="w",padx=16,pady=18)
add_nav("Dashboard","⌂   Dashboard",show_dashboard); add_nav("Book Ticket","▣   Book Ticket",show_booking); add_nav("Bookings","▤   Bookings",show_bookings); add_nav("Reports","▥   Reports",show_reports); add_nav("Settings","⚙   Settings",show_settings)
show_booking(); root.mainloop()
